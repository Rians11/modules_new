"""Generate the Cinema Management System Excel workbook (full, mirroring the lecturer's example)."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

TITLE_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FILL = PatternFill("solid", fgColor="4472C4")
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
WHITE = Font(color="FFFFFF", bold=True)
BOLD = Font(bold=True)
thin = Side(style="thin", color="B0B0B0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = WHITE
        cell.alignment = CENTER
        cell.border = BORDER


def box(ws, r1, r2, ncols):
    for r in range(r1, r2 + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            if not cell.alignment.wrap_text:
                cell.alignment = WRAP


wb = openpyxl.Workbook()

# ---- Sheet 1 : Use Cases ----------------------------------------------------
ws = wb.active
ws.title = "Use Cases"
ws.merge_cells("A1:C1")
t = ws["A1"]
t.value = "Cinema Management System"
t.fill = TITLE_FILL
t.font = Font(color="FFFFFF", bold=True, size=14)
t.alignment = CENTER
ws.row_dimensions[1].height = 24
ws.append(["No", "Use Case", "Actor"])
style_header(ws, 2, 3)

use_cases = [
    ("Add Movie", "Manager"),
    ("Update Movie", "Manager"),
    ("Delete Movie", "Manager"),
    ("Search Movie by Title, Genre, Language", "Manager, Cashier"),
    ("Add Genre", "Manager"),
    ("Update Genre", "Manager"),
    ("Delete Genre", "Manager"),
    ("Search Genre by Name", "Manager"),
    ("Add Cinema Hall", "Manager"),
    ("Update Cinema Hall", "Manager"),
    ("Delete Cinema Hall", "Manager"),
    ("Search Cinema Hall by Name, Type, Capacity", "Manager"),
    ("Add Showtime", "Manager"),
    ("Update Showtime", "Manager"),
    ("Delete Showtime", "Manager"),
    ("Search Showtime by Date, Movie, Hall, Time", "Manager, Cashier"),
    ("Add Customer", "Cashier"),
    ("Update Customer", "Cashier"),
    ("Delete Customer", "Manager"),
    ("Search Customer by Id, Name, Phone", "Cashier"),
    ("Add Staff", "Manager"),
    ("Update Staff", "Manager"),
    ("Delete Staff", "Manager"),
    ("Search Staff by Id, Name, Role", "Manager"),
    ("Add Booking", "Cashier, Customer (secondary)"),
    ("Update Booking", "Cashier"),
    ("Delete Booking", "Cashier"),
    ("Search Booking by Id, Customer, Date, Showtime", "Cashier, Manager"),
    ("Add Ticket", "Cashier"),
    ("Update Ticket", "Cashier"),
    ("Delete Ticket", "Cashier"),
    ("Search Ticket by Code, Showtime, Customer", "Cashier"),
    ("Add Payment", "Cashier"),
    ("Update Payment", "Cashier"),
    ("Delete Payment", "Manager"),
    ("Search Payment by Id, Date, Method, Customer", "Cashier, Manager"),
    ("Add Snack", "Manager"),
    ("Update Snack", "Manager"),
    ("Delete Snack", "Manager"),
    ("Search Snack by Name, Category", "Cashier, Manager"),
    ("Add Promotion", "Manager"),
    ("Update Promotion", "Manager"),
    ("Delete Promotion", "Manager"),
    ("Search Promotion by Code, Period", "Cashier, Manager"),
    ("Post Showtime", "Manager"),
    ("Unpost Showtime", "Manager"),
    ("Select Seat", "Cashier, Customer (secondary)"),
    ("Apply Promotion to Booking", "Cashier"),
    ("Generate Ticket", "Cashier"),
    ("Validate Ticket at Entrance", "Cashier"),
    ("Issue Refund", "Cashier, Manager"),
    ("Generate Sales Report", "Manager"),
    ("Generate Hall Occupancy Report", "Manager"),
    ("Generate Daily Revenue Report", "Manager"),
    ("Manage User Roles", "Manager"),
    ("Login", "Manager, Cashier"),
    ("Logout", "Manager, Cashier"),
]
for i, (uc, actor) in enumerate(use_cases, start=1):
    ws.append([i, uc, actor])
    ws.cell(row=2 + i, column=1).alignment = CENTER
box(ws, 2, 2 + len(use_cases), 3)
ws.column_dimensions["A"].width = 6
ws.column_dimensions["B"].width = 48
ws.column_dimensions["C"].width = 30
ws.freeze_panes = "A3"

# ---- Sheet 2 : Use Case Specifications -------------------------------------
ws2 = wb.create_sheet("Use Case Specifications")
specs = [
    ("UC-25", "Add Booking", "Cashier (primary), Customer (secondary)",
     "Creates a booking for a showtime: selects seats, applies any promotion, takes payment and issues the ticket.",
     "Cashier is logged in; a posted showtime with a free seat exists.",
     "A booking, ticket and payment are saved; the seats are marked sold.",
     "1. Select a posted showtime. 2. Select seats (include Select Seat). 3. Enter/select the customer. "
     "4. (optional) Apply a promotion (extend Apply Promotion). 5. Show total and take payment (include "
     "Add Payment). 6. Save booking and issue ticket (include Generate Ticket); mark seats sold.",
     "2a. Seat already sold -> choose another seat. 5a. Payment fails -> release seats, no ticket."),
    ("UC-50", "Validate Ticket at Entrance", "Cashier",
     "Checks a presented ticket is valid for the current showtime and marks it used.",
     "Cashier is logged in; the ticket exists.",
     "Ticket marked Used; entry granted or refused.",
     "1. Scan/enter the ticket code. 2. Find ticket and showtime. 3. Check valid, right showtime, not used. "
     "4. Mark Used and show 'Access granted' with the seat number.",
     "3a. Already used / wrong showtime -> 'Invalid ticket', refuse entry."),
    ("UC-52", "Generate Sales Report", "Manager",
     "Produces a report of tickets sold and revenue for a chosen period.",
     "Manager is logged in.",
     "A report is displayed/printed.",
     "1. Choose the period (date range). 2. System totals tickets sold and revenue. 3. Display report with grand total.",
     "1a. No data in period -> report shows zero totals."),
    ("UC-01", "Add Movie", "Manager",
     "Adds a new movie to the catalogue (template for all Add/Update/Delete/Search use cases).",
     "Manager is logged in.",
     "A new movie record is stored.",
     "1. Open the Movie form. 2. Enter title, genre, duration, language, rating. 3. Save. 4. Validate and store.",
     "3a. Duplicate / missing field -> error, not saved."),
]
fields = ["Use Case ID", "Use Case Name", "Actor(s)", "Description",
          "Preconditions", "Postconditions", "Main Flow", "Alternative Flows"]
r = 1
for uc_id, name, actors, desc, pre, post, main, alt in specs:
    ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    h = ws2.cell(row=r, column=1, value=f"{uc_id}  —  {name}")
    h.fill = HEAD_FILL
    h.font = WHITE
    h.alignment = Alignment(horizontal="left", vertical="center")
    ws2.cell(row=r, column=2).fill = HEAD_FILL
    r += 1
    for fname, val in zip(fields, [uc_id, name, actors, desc, pre, post, main, alt]):
        fc = ws2.cell(row=r, column=1, value=fname)
        fc.fill = SUB_FILL
        fc.font = BOLD
        fc.alignment = WRAP
        ws2.cell(row=r, column=2, value=val).alignment = WRAP
        r += 1
    r += 1
box(ws2, 1, r, 2)
ws2.column_dimensions["A"].width = 22
ws2.column_dimensions["B"].width = 95
note = ws2.cell(row=r + 1, column=1,
                value="Replicate the CRUD template for every Add/Update/Delete/Search use case of Movie, Genre, "
                      "Hall, Showtime, Customer, Staff, Booking, Ticket, Payment, Snack, Promotion. Write full "
                      "specs for Post/Unpost Showtime, Select Seat, Apply Promotion, Add Payment, Generate Ticket, "
                      "Issue Refund, the three reports, Manage User Roles, Login, Logout.")
note.font = Font(italic=True, color="808080")
note.alignment = WRAP
ws2.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=2)

# ---- Sheet 3 : Database Design ---------------------------------------------
ws3 = wb.create_sheet("Database Design")
tables = {
    "staff": [("staff_id", "INT (auto)", "PK"), ("full_name", "VARCHAR(100)", ""),
              ("username", "VARCHAR(50)", ""), ("password", "VARCHAR(255)", ""),
              ("role", "ENUM('MANAGER','CASHIER')", "")],
    "customer": [("customer_id", "INT (auto)", "PK"), ("full_name", "VARCHAR(100)", ""),
                 ("phone", "VARCHAR(20)", ""), ("email", "VARCHAR(100)", "")],
    "genre": [("genre_id", "INT (auto)", "PK"), ("name", "VARCHAR(50)", "")],
    "movie": [("movie_id", "INT (auto)", "PK"), ("title", "VARCHAR(120)", ""),
              ("genre_id", "INT", "FK -> genre"), ("duration", "INT", ""),
              ("language", "VARCHAR(40)", ""), ("rating", "VARCHAR(10)", "")],
    "cinema_hall": [("hall_id", "INT (auto)", "PK"), ("hall_name", "VARCHAR(60)", ""),
                    ("hall_type", "VARCHAR(20)", ""), ("capacity", "INT", "")],
    "showtime": [("show_id", "INT (auto)", "PK"), ("movie_id", "INT", "FK -> movie"),
                 ("hall_id", "INT", "FK -> cinema_hall"), ("show_date", "DATE", ""),
                 ("show_time", "TIME", ""), ("ticket_price", "DECIMAL(8,2)", ""),
                 ("status", "ENUM('Posted','Unposted')", "")],
    "booking": [("booking_id", "INT (auto)", "PK"), ("customer_id", "INT", "FK -> customer"),
                ("show_id", "INT", "FK -> showtime"), ("promotion_id", "INT", "FK -> promotion (null)"),
                ("booking_datetime", "DATETIME", ""), ("status", "ENUM('Confirmed','Cancelled')", ""),
                ("total_amount", "DECIMAL(8,2)", "")],
    "ticket": [("ticket_id", "INT (auto)", "PK"), ("booking_id", "INT", "FK -> booking"),
               ("seat_number", "VARCHAR(8)", ""), ("price", "DECIMAL(8,2)", ""),
               ("status", "ENUM('Valid','Used','Refunded')", "")],
    "payment": [("payment_id", "INT (auto)", "PK"), ("booking_id", "INT", "FK -> booking"),
                ("amount", "DECIMAL(8,2)", ""), ("method", "ENUM('Card','Mobile','Cash')", ""),
                ("payment_datetime", "DATETIME", ""), ("status", "ENUM('Paid','Refunded')", "")],
    "snack": [("snack_id", "INT (auto)", "PK"), ("name", "VARCHAR(50)", ""),
              ("category", "VARCHAR(30)", ""), ("price", "DECIMAL(8,2)", "")],
    "snack_order": [("order_id", "INT (auto)", "PK"), ("booking_id", "INT", "FK -> booking"),
                    ("snack_id", "INT", "FK -> snack"), ("quantity", "INT", ""),
                    ("line_total", "DECIMAL(8,2)", "")],
    "promotion": [("promotion_id", "INT (auto)", "PK"), ("code", "VARCHAR(20)", ""),
                  ("description", "VARCHAR(100)", ""), ("discount_percent", "DECIMAL(5,2)", ""),
                  ("start_date", "DATE", ""), ("end_date", "DATE", "")],
}
r = 1
for tname, cols in tables.items():
    ws3.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    th = ws3.cell(row=r, column=1, value=tname)
    th.fill = TITLE_FILL
    th.font = Font(color="FFFFFF", bold=True)
    r += 1
    for j, head in enumerate(["Attribute", "Data Type", "Key"], start=1):
        c = ws3.cell(row=r, column=j, value=head)
        c.fill = SUB_FILL
        c.font = BOLD
        c.border = BORDER
    r += 1
    for attr, dtype, key in cols:
        ws3.cell(row=r, column=1, value=attr).border = BORDER
        ws3.cell(row=r, column=2, value=dtype).border = BORDER
        ws3.cell(row=r, column=3, value=key).border = BORDER
        r += 1
    r += 1
ws3.column_dimensions["A"].width = 20
ws3.column_dimensions["B"].width = 30
ws3.column_dimensions["C"].width = 24

out = "Cinema_Management_System.xlsx"
wb.save(out)
print("Saved", out)
